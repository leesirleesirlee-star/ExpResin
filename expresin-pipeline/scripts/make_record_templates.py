"""生成三套标准原始数据记录表（空白模板 + 示例填写）到 materials/record-templates/。

设计依据：docs/record-template-design.md
  - 一个实验一个 .xlsx，单数据 sheet（sheet 名 Record）；
  - 顶部实验信息区：每行单格、供人读与追溯、不参与自动识别；
  - 数据块：块标题行（写明离子名；多室电化学为「【室名】【离子名】测量数据」）
    + 单行表头（英文列名 + 括号单位）+ 连续数据行；
  - 多离子（及多室）= 多块垂直堆叠，块间保留 3 行空行（L1 以连续 3 空行判定块边界）；
  - 电化学模板含 Initial Concentration (mg/L) 行级初始浓度列：同一数据块内
    多个初始浓度（多轮实验）按轮次连排，块内不留空行；
  - 表头行正下方必须紧跟数据（L1 表头判定要求「下一行非空」）。

生成后立即用 L1 结构解析自检：
  - 示例文件应解析出预期数量的数据块（region）；
  - 空白模板应解析出 0 个数据块（表头下方无数据）。

用法:
    python expresin-pipeline/scripts/make_record_templates.py [--out DIR]
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "expresin-pipeline" / "src"))

from l1_structure import parse_workbook  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

CONFIG_DIR = ROOT / "expresin-pipeline" / "config" / "record_templates"
DEFAULT_OUT = ROOT / "materials" / "record-templates"

# 文件名前缀（仓库规范：文件名英文）
FILE_PREFIX = {
    "batch": "Batch_Record_Template_v1",
    "column": "Column_Record_Template_v1",
    "electrochemical": "Electrochemical_Record_Template_v1",
    # 电化学分支（两室 / 三室 / 双极膜）
    "electrochemical_two_chamber": "Electrochemical_TwoChamber_Record_Template_v1",
    "electrochemical_three_chamber": "Electrochemical_ThreeChamber_Record_Template_v1",
    "electrochemical_bipolar": "Electrochemical_BipolarMembrane_Record_Template_v1",
}

# 信息区 / 数据块 / 说明区之间的空行数（L1 以连续 3 空行判定块边界，不得小于 3）
SECTION_GAP = 3
BLANK_PLACEHOLDER = "______"

COL_WIDTHS = {"A": 34, "B": 26, "C": 18, "D": 20, "E": 18}
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(color="808080", size=10)
SECTION_FONT = Font(bold=True, size=11)
HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
LEFT = Alignment(horizontal="left")


def build_sheet(ws, spec: dict, filled: bool) -> None:
    """按规范往工作表写内容；filled=True 时写入示例数据（用于示例文件）。"""
    row = 1
    ws.cell(row=row, column=1, value=spec["title"]).font = TITLE_FONT
    row += 2
    ws.cell(row=row, column=1, value=spec["subtitle"]).font = SUBTITLE_FONT
    row += 2

    ws.cell(row=row, column=1, value="实验信息（每行一格，供人读与追溯，不参与自动识别）").font = SECTION_FONT
    row += 1
    example_values = (spec.get("example") or {}).get("info_values") or []
    for index, field in enumerate(spec["info_section"]["fields"]):
        if filled and index < len(example_values):
            text = f"{field['label']}：{example_values[index]}"
        else:
            text = f"{field['label']}：{BLANK_PLACEHOLDER}"
        ws.cell(row=row, column=1, value=text).alignment = LEFT
        row += 1

    row += SECTION_GAP

    blocks = ((spec.get("example") or {}).get("blocks") or []) if filled else []
    if not blocks:
        # 空白模板：优先使用 spec.blank_blocks（多室电化学每室一个占位块），
        # 否则退化为单一 block_template 占位块。
        titles = spec.get("blank_blocks") or [spec["block_template"]["title"]]
        blocks = [{"title": title, "rows": []} for title in titles]

    for block_index, block in enumerate(blocks):
        if block_index > 0:
            row += SECTION_GAP
        ws.cell(row=row, column=1, value=block["title"]).font = SECTION_FONT
        row += 1
        for col_index, column in enumerate(spec["data_columns"], start=1):
            cell = ws.cell(row=row, column=col_index, value=column["header"])
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
        row += 1
        for data_row in block.get("rows") or []:
            for col_index, value in enumerate(data_row, start=1):
                ws.cell(row=row, column=col_index, value=value)
            row += 1

    row += SECTION_GAP
    ws.cell(row=row, column=1, value="填写说明").font = SECTION_FONT
    row += 1
    for rule in spec.get("fill_rules") or []:
        ws.cell(row=row, column=1, value=rule).alignment = LEFT
        row += 1

    for col_letter, width in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width


def self_check(path: Path, expect_regions: int) -> bool:
    """用 L1 结构解析验证生成结果（示例块数与预期一致 / 空白表无数据块）。"""
    structure = parse_workbook(path)
    sheet = structure.sheets[0]
    ok = sheet.region_count == expect_regions
    print(f"  [self-check] {path.name}: regions={sheet.region_count} (期望 {expect_regions}) {'OK' if ok else 'MISMATCH'}")
    for region in sheet.regions:
        print(
            f"      - {region.region_id} title={region.title!r} "
            f"cols={region.col_count} rows={region.row_count} kind={region.region_kind}"
        )
    return ok


def main() -> int:
    parser = argparse.ArgumentParser(description="生成标准原始数据记录表（空白 + 示例）")
    parser.add_argument("--out", default=str(DEFAULT_OUT), help="输出目录（默认 materials/record-templates）")
    args = parser.parse_args()
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    ok = True
    for config_file in sorted(CONFIG_DIR.glob("*_record_v1.json")):
        spec = json.loads(config_file.read_text(encoding="utf-8"))
        prefix = FILE_PREFIX[spec["experiment_type"]]
        expected_regions = len((spec.get("example") or {}).get("blocks") or [])

        for filled, suffix in ((False, ""), (True, "_example_filled")):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = spec["sheet_name"]
            build_sheet(sheet, spec, filled=filled)
            path = out_dir / f"{prefix}{suffix}.xlsx"
            workbook.save(path)
            print(f"generated: {path.relative_to(ROOT)}")
            ok = self_check(path, expected_regions if filled else 0) and ok

    print("==== all checks passed ====" if ok else "==== checks FAILED ====")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
