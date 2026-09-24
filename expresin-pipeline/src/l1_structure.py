"""L1 确定性结构解析层 (Structure-First Parsing) — v2 多数据块版本.

目的
----
把"非标准、多 Sheet、多数据块、多级表头、单位混排"的真实实验 Excel，
还原成**带坐标的结构化描述**，作为后续 L2~L5 语义映射与证据校验的输入。

v1 -> v2 的关键修正
------------------
真实实验表格的一个 Sheet 内部往往包含**多个独立数据块**，例如：
    R1  实验条件描述（自由文本）
    R2  [标定曲线块] 表头:  SO4 | area | ppm |  | Cl | area | ppm
    R3~R8              标定曲线数据
    R9~R11             空行分隔
    R12 区块标题:      SO4
    R13 [样品数据块] 表头: BV | V | number | area | calculated ppm | 稀释100倍(mol/L) | mg/L
    R14~R35            样品测量数据
若只识别"一个表头"，就会丢掉标定曲线或样品数据。v2 因此实现：
    1. 扫描并定位所有表头行；
    2. 相邻表头行聚成一个多级表头块 (header block)；
    3. 每个表头块对应一个 DataRegion，向下延伸至下一个表头块或连续空行；
    4. 在每个 Region 内按"表头与数据均空的列"切分列分组 (ColumnBlock)。

设计原则
--------
1. 本层完全不依赖 LLM：结构解析必须确定、可复现、可测试。
2. 输出"结构化摘要"而非原始文本，避免整表塞入模型导致行列关系丢失。
3. 每个结论都带证据（行号、列号、样本值），便于后续置信度评估。
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# 基础判定工具
# --------------------------------------------------------------------------- #

_NUM_RE = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")
_UNIT_PAREN_RE = re.compile(r"[\(\[]\s*([^\)\]]+?)\s*[\)\]]")

# 全角 -> 半角（表格里常见全角括号/逗号/冒号）
_FULLWIDTH_MAP = str.maketrans(
    {
        "（": "(",
        "）": ")",
        "［": "[",
        "］": "]",
        "【": "[",
        "】": "]",
        "，": ",",
        "：": ":",
        "　": " ",
    }
)

# 常见实验单位 -> 规范写法
_KNOWN_UNITS: dict[str, str] = {
    "ppm": "ppm", "ppb": "ppb", "mg/l": "mg/L", "mgl": "mg/L",
    "ug/l": "ug/L", "µg/l": "ug/L", "g/l": "g/L", "mg/ml": "mg/mL",
    "mol/l": "mol/L", "mmol/l": "mmol/L", "umol/l": "umol/L",
    "ml": "mL", "l": "L", "ul": "uL",
    "mg": "mg", "g": "g", "kg": "kg", "ug": "ug",
    "min": "min", "mins": "min", "minute": "min", "minutes": "min",
    "h": "h", "hr": "h", "hrs": "h", "hour": "h", "hours": "h",
    "s": "s", "sec": "s", "second": "s", "seconds": "s",
    "day": "day", "days": "day", "d": "day",
    "c": "degC", "℃": "degC", "°c": "degC", "k": "K",
    "%": "%", "v": "V", "mv": "mV", "a": "A", "ma": "mA",
    "ph": "pH", "rpm": "rpm", "bv": "BV", "bar": "bar",
    "kpa": "kPa", "mpa": "MPa", "nm": "nm", "um": "um",
    "cm": "cm", "mm": "mm",
}


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def is_number(value: Any) -> bool:
    """判断单元格值是否为数值（含可解析的数值字符串）。"""
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        return bool(_NUM_RE.match(value.strip()))
    return False


def to_float(value: Any) -> float | None:
    if is_number(value):
        try:
            return float(value)
        except (TypeError, ValueError):
            return None
    return None


def extract_unit(header: str) -> str | None:
    """从列名中抽取单位。

    抽取顺序（从严到宽）：
    1. 括号内单位，如 "稀释100倍(mol/L)" -> mol/L；
    2. 整个表头就是单位，如 "mg/L"、"ppm"；
    3. 由空格/逗号/斜杠分隔的单位 token，如 "time,min" -> min。

    注意：**单字母表头一律不视作单位**。因为在此实验数据中 "V" 表示体积(volume)、
    "A"/"B" 常为分组标签，把单字母当单位会产生大量假阳性（曾把 V 误读成伏特）。
    """
    if not header:
        return None

    for m in _UNIT_PAREN_RE.finditer(header):
        inner = (m.group(1) or "").strip().lower()
        if inner in _KNOWN_UNITS:
            return _KNOWN_UNITS[inner]

    whole = header.strip().lower()
    if len(whole) >= 2 and whole in _KNOWN_UNITS:
        return _KNOWN_UNITS[whole]

    for token in re.split(r"[\s,;/]+", header.strip().lower()):
        token = token.strip("()[].")
        if len(token) >= 2 and token in _KNOWN_UNITS:
            return _KNOWN_UNITS[token]
    return None


def normalize_name(header: str) -> str:
    """列名归一化：统一括号/大小写/空白，剥离单位，便于检索与比对。"""
    if not header:
        return ""
    s = header.strip().translate(_FULLWIDTH_MAP)
    s = _UNIT_PAREN_RE.sub(" ", s)
    s = re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff+\-]", " ", s)
    s = re.sub(r"\s+", " ", s).strip().lower()
    return s


def classify_values(values: list[Any]) -> str:
    nums = texts = 0
    for v in values:
        if is_blank(v):
            continue
        if is_number(v):
            nums += 1
        else:
            texts += 1
    if nums and texts:
        return "mixed"
    if nums:
        return "number"
    if texts:
        return "text"
    return "empty"


# --------------------------------------------------------------------------- #
# 数据结构
# --------------------------------------------------------------------------- #


@dataclass
class ColumnProfile:
    """单列的结构化画像。"""

    col_index: int
    col_letter: str
    raw_header: str
    normalized: str
    unit: str | None
    header_path: list[str]
    value_type: str
    non_null: int
    fill_ratio: float
    sample_values: list[Any] = field(default_factory=list)
    numeric_range: list[float] | None = None


@dataclass
class ColumnBlock:
    """Region 内按空列切分出的列分组（例如 SO4 组 / Cl 组）。"""

    block_id: str
    col_start: int
    col_end: int
    col_letters: str
    group_hint: str | None
    columns: list[ColumnProfile] = field(default_factory=list)


@dataclass
class DataRegion:
    """一个 Sheet 内的一个数据块。"""

    region_id: str
    region_index: int
    region_kind: str
    title: str | None
    header_rows: list[int]
    data_start_row: int
    data_end_row: int
    row_count: int
    col_count: int
    col_blocks: list[ColumnBlock] = field(default_factory=list)


@dataclass
class SheetStructure:
    sheet_name: str
    sheet_index: int
    dimensions: str
    max_row: int
    max_col: int
    merged_ranges: list[str] = field(default_factory=list)
    region_count: int = 0
    regions: list[DataRegion] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class WorkbookStructure:
    file_name: str
    file_path: str
    sheet_count: int
    sheets: list[SheetStructure] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# 解析算法
# --------------------------------------------------------------------------- #


def _row_profile(ws, row: int, max_col: int) -> tuple[int, int, int]:
    """返回 (文本单元格数, 数值单元格数, 非空单元格数)。"""
    text = num = nonempty = 0
    for c in range(1, max_col + 1):
        v = ws.cell(row=row, column=c).value
        if is_blank(v):
            continue
        nonempty += 1
        if is_number(v):
            num += 1
        else:
            text += 1
    return text, num, nonempty


def find_header_blocks(ws, max_col: int) -> list[list[int]]:
    """定位所有表头行并聚成多级表头块。

    表头行判定：非空 >= 2；文本单元格数 >= 数值单元格数（数值为主的是数据行）；
    下一行存在非空内容。相邻表头行合并为一个多级表头块。
    """
    max_row = ws.max_row or 0
    candidates: list[int] = []
    for r in range(1, max_row + 1):
        text, num, nonempty = _row_profile(ws, r, max_col)
        if nonempty == 0 or text < 2:
            continue
        if num > text:
            continue
        if r < max_row:
            _, _, next_nonempty = _row_profile(ws, r + 1, max_col)
            if next_nonempty == 0:
                continue
        candidates.append(r)

    blocks: list[list[int]] = []
    current: list[int] = []
    for r in candidates:
        if current and r == current[-1] + 1:
            current.append(r)
        else:
            if current:
                blocks.append(current)
            current = [r]
    if current:
        blocks.append(current)
    return blocks


def build_merged_lookup(ws) -> dict[tuple[int, int], tuple[int, int]]:
    """合并单元格映射：任意坐标 -> 合并区左上角坐标。"""
    lookup: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                lookup[(row, col)] = (rng.min_row, rng.min_col)
    return lookup


def _header_value(ws, row: int, col: int, merged: dict[tuple[int, int], tuple[int, int]]) -> str:
    key = (row, col)
    if key in merged:
        mr, mc = merged[key]
        v = ws.cell(row=mr, column=mc).value
    else:
        v = ws.cell(row=row, column=col).value
    return "" if is_blank(v) else str(v).strip()


def detect_region_end(
    ws, start_row: int, max_col: int, next_header_row: int | None, blank_tolerance: int = 3
) -> int:
    """从 data_start 起向下延伸数据区，返回数据区结束行。"""
    max_row = ws.max_row or 0
    limit = (next_header_row - 1) if next_header_row else max_row
    end = start_row - 1
    blank_run = 0
    for r in range(start_row, limit + 1):
        nonempty = 0
        for c in range(1, max_col + 1):
            if not is_blank(ws.cell(row=r, column=c).value):
                nonempty += 1
        if nonempty == 0:
            blank_run += 1
            if blank_run >= blank_tolerance:
                break
        else:
            blank_run = 0
            end = r
    return end


def find_title(ws, header_rows: list[int], max_col: int) -> str | None:
    """取表头块上方最近的非空文本，通常是实验条件描述或区块标题。"""
    for r in range(header_rows[0] - 1, 0, -1):
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if not is_blank(v):
                return str(v).strip()
    return None


def build_column_profile(
    ws, col: int, header_rows: list[int], data_rows: list[int], merged: dict
) -> ColumnProfile:
    raw_header = _header_value(ws, header_rows[-1], col, merged)
    path = []
    for r in header_rows:
        v = _header_value(ws, r, col, merged)
        if v and v not in path:
            path.append(v)

    values = [ws.cell(row=r, column=col).value for r in data_rows]
    clean = [v for v in values if not is_blank(v)]
    numbers = [f for f in (to_float(v) for v in clean) if f is not None]

    return ColumnProfile(
        col_index=col,
        col_letter=get_column_letter(col),
        raw_header=raw_header,
        normalized=normalize_name(raw_header),
        unit=extract_unit(raw_header),
        header_path=path,
        value_type=classify_values(values),
        non_null=len(clean),
        fill_ratio=round(len(clean) / len(data_rows), 3) if data_rows else 0.0,
        sample_values=clean[:5],
        numeric_range=[min(numbers), max(numbers)] if numbers else None,
    )


def split_column_blocks(
    ws, header_rows: list[int], data_rows: list[int], max_col: int, merged: dict, region_id: str
) -> list[ColumnBlock]:
    """按"表头与数据均空"的列切分列分组。"""
    is_active: list[bool] = []
    for c in range(1, max_col + 1):
        header_active = any(_header_value(ws, r, c, merged) for r in header_rows)
        data_active = any(not is_blank(ws.cell(row=r, column=c).value) for r in data_rows)
        is_active.append(header_active or data_active)

    spans: list[tuple[int, int]] = []
    start: int | None = None
    for idx, active in enumerate(is_active, start=1):
        if active:
            if start is None:
                start = idx
        else:
            if start is not None:
                spans.append((start, idx - 1))
                start = None
    if start is not None:
        spans.append((start, len(is_active)))

    blocks: list[ColumnBlock] = []
    for bi, (cs, ce) in enumerate(spans, start=1):
        columns = [build_column_profile(ws, c, header_rows, data_rows, merged) for c in range(cs, ce + 1)]
        hint = columns[0].raw_header if columns and columns[0].raw_header else None
        letters = f"{get_column_letter(cs)}-{get_column_letter(ce)}" if ce > cs else get_column_letter(cs)
        blocks.append(
            ColumnBlock(
                block_id=f"{region_id}/B{bi}",
                col_start=cs,
                col_end=ce,
                col_letters=letters,
                group_hint=hint,
                columns=columns,
            )
        )
    return blocks


def guess_region_kind(columns: list[ColumnProfile]) -> str:
    """启发式推断数据块语义类型，辅助后续映射。"""
    names = " ".join(c.normalized for c in columns)
    if "bed volume" in names or re.search(r"\bbv\b", names):
        return "sample_data"
    if "capacity" in names or "mg/g" in names:
        return "capacity"
    if "time" in names and "ph" in names:
        return "ph_curve"
    if re.search(r"\barea\b", names) and re.search(r"\bppm\b", names):
        return "calibration"
    if "time" in names:
        return "kinetic"
    return "unknown"


def parse_sheet(ws, sheet_index: int) -> SheetStructure:
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    sheet = SheetStructure(
        sheet_name=ws.title,
        sheet_index=sheet_index,
        dimensions=ws.dimensions,
        max_row=max_row,
        max_col=max_col,
        merged_ranges=[str(r) for r in ws.merged_cells.ranges],
    )

    if max_row < 2 or max_col < 1:
        sheet.warnings.append("Sheet 过小，未识别数据区")
        return sheet

    if ws.title != ws.title.strip():
        sheet.warnings.append("Sheet 名包含首尾空格")

    merged = build_merged_lookup(ws)
    header_blocks = find_header_blocks(ws, max_col)

    regions: list[DataRegion] = []
    for bi, header_rows in enumerate(header_blocks):
        data_start = header_rows[-1] + 1
        next_header = header_blocks[bi + 1][0] if bi + 1 < len(header_blocks) else None
        data_end = detect_region_end(ws, data_start, max_col, next_header)
        if data_end < data_start:
            continue

        data_rows = list(range(data_start, data_end + 1))
        region_id = f"{ws.title}#R{header_rows[0]}-{header_rows[-1]}"
        col_blocks = split_column_blocks(ws, header_rows, data_rows, max_col, merged, region_id)
        flat_columns = [c for b in col_blocks for c in b.columns]

        regions.append(
            DataRegion(
                region_id=region_id,
                region_index=len(regions) + 1,
                region_kind=guess_region_kind(flat_columns),
                title=find_title(ws, header_rows, max_col),
                header_rows=header_rows,
                data_start_row=data_start,
                data_end_row=data_end,
                row_count=len(data_rows),
                col_count=len(flat_columns),
                col_blocks=col_blocks,
            )
        )

    if not regions:
        sheet.warnings.append("未识别到有效数据块")

    sheet.regions = regions
    sheet.region_count = len(regions)
    return sheet


def parse_workbook(path: str | Path) -> WorkbookStructure:
    """解析整个工作簿为结构化描述。"""
    p = Path(path)
    wb = openpyxl.load_workbook(p, data_only=True)
    try:
        sheets = [parse_sheet(ws, i) for i, ws in enumerate(wb.worksheets)]
    finally:
        wb.close()
    return WorkbookStructure(
        file_name=p.name,
        file_path=str(p),
        sheet_count=len(sheets),
        sheets=sheets,
    )


def to_dict(structure: WorkbookStructure) -> dict:
    return asdict(structure)


def summarize(structure: WorkbookStructure) -> str:
    """生成人类可读摘要，用于命令行快速检查。"""
    lines = [f"FILE: {structure.file_name}", f"sheets: {structure.sheet_count}"]
    for s in structure.sheets:
        lines.append(f"  [{s.sheet_name}] regions={s.region_count}")
        for r in s.regions:
            title = (r.title or "")[:46]
            lines.append(
                f"    - R{r.region_index} kind={r.region_kind} header={r.header_rows} "
                f"data={r.data_start_row}-{r.data_end_row} ({r.row_count}行 x {r.col_count}列) title={title}"
            )
            for b in r.col_blocks:
                desc = " | ".join(f"{c.col_letter}:{c.raw_header or '<空>'}" for c in b.columns[:8])
                lines.append(f"        块{b.col_letters} [{b.group_hint}] {desc}")
        if s.warnings:
            lines.append(f"    ! {'; '.join(s.warnings)}")
    return "\n".join(lines)


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("usage: python l1_structure.py <excel path>")
        raise SystemExit(1)
    print(summarize(parse_workbook(sys.argv[1])))