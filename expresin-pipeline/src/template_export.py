"""ExpResin 模板 Excel 导出（template_export_v1）。

把 template_fill 的结果写成 xlsx：**每个目标离子一个工作表**，
每张表按模板的 A/B/C 三段结构，附数据录入规则与口径说明（Provenance）。

约定：
  - 表头语言与模板一致（英文），忠于老师提供模板；
  - 缺失值留空（不写 "NA"，避免与「不适用」混淆），C 区无法计算的原因写在 Note 列；
  - 不写任何未经计算或推断的数字。
"""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXPORT_VERSION = "template_export_v1"

_INVALID_SHEET_CHARS = set('[]:*?/\\')

_TITLE_FONT = Font(bold=True, size=14)
_SUBTITLE_FONT = Font(size=9, color="808080")
_SECTION_FONT = Font(bold=True, size=12)
_HEAD_FONT = Font(bold=True)
_HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
_MUTED_FONT = Font(size=9, color="808080")


def _safe_sheet_name(name: str, used: set[str]) -> str:
    """清理非法字符并保证工作表名唯一（Excel 限制 31 字符）。"""
    cleaned = "".join("_" if ch in _INVALID_SHEET_CHARS else ch for ch in str(name or "Sheet")).strip()
    cleaned = (cleaned or "Sheet")[:31]
    candidate = cleaned
    index = 2
    while candidate.lower() in used:
        suffix = f"_{index}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate.lower())
    return candidate


def _as_cell(value):
    """openpyxl 只接受标量；其余转字符串。"""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _write_headers(ws, row: int, labels: list[str]) -> int:
    for index, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=index, value=label)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
    return row + 1


def _write_sheet(ws, data: dict, sheet: dict) -> None:
    sections = sheet.get("sections") or {}
    row = 1

    ws.cell(row=row, column=1, value=data.get("template_title") or "EXPERIMENT").font = _TITLE_FONT
    row += 1
    subtitle = (
        f"Experiment {data.get('experiment_id')} · "
        f"analyte {sheet.get('analyte') or '—'} · "
        f"template {data.get('template_id')} v{data.get('template_version')} · "
        f"generated {data.get('generated_at')}"
    )
    ws.cell(row=row, column=1, value=subtitle).font = _SUBTITLE_FONT
    row += 2

    # ---- A. Experiment Information ----
    section_a = sections.get("A") or {}
    ws.cell(row=row, column=1, value=f"A. {section_a.get('label')}").font = _SECTION_FONT
    row += 1
    row = _write_headers(ws, row, ["Field", "Value", "Unit"])
    for item in section_a.get("rows") or []:
        ws.cell(row=row, column=1, value=item.get("label"))
        ws.cell(row=row, column=2, value=_as_cell(item.get("value")))
        ws.cell(row=row, column=3, value=item.get("unit") or None)
        row += 1
    row += 1

    # ---- B. Experimental Data ----
    section_b = sections.get("B") or {}
    ws.cell(row=row, column=1, value=f"B. {section_b.get('label')}").font = _SECTION_FONT
    row += 1
    columns = section_b.get("columns") or []
    labels = []
    for col in columns:
        unit = col.get("unit")
        labels.append(f"{col.get('label')} ({unit})" if unit else str(col.get("label")))
    row = _write_headers(ws, row, labels)
    for data_row in section_b.get("rows") or []:
        for index, value in enumerate(data_row, start=1):
            ws.cell(row=row, column=index, value=_as_cell(value))
        row += 1
    if not (section_b.get("rows") or []):
        ws.cell(row=row, column=1, value="(no rows identified)").font = _MUTED_FONT
        row += 1
    row += 1

    # ---- C. Derived ----
    section_c = sections.get("C") or {}
    ws.cell(row=row, column=1, value=f"C. {section_c.get('label')}").font = _SECTION_FONT
    row += 1
    row = _write_headers(ws, row, ["Parameter", "Value", "Unit", "Note"])
    for item in section_c.get("rows") or []:
        ws.cell(row=row, column=1, value=item.get("label"))
        ws.cell(row=row, column=2, value=_as_cell(item.get("value")))
        ws.cell(row=row, column=3, value=item.get("unit") or None)
        note = item.get("reason") or item.get("method") or ""
        ws.cell(row=row, column=4, value=note)
        if item.get("value") is None and item.get("reason"):
            ws.cell(row=row, column=2).font = _MUTED_FONT
        row += 1
    row += 1

    # ---- 数据录入规则 ----
    rules = data.get("rules") or []
    if rules:
        ws.cell(row=row, column=1, value="Data Entry Rules").font = _SECTION_FONT
        row += 1
        for rule in rules:
            ws.cell(row=row, column=1, value=f"* {rule}")
            row += 1
        row += 1

    # ---- 口径与来源 ----
    provenance = data.get("provenance") or {}
    ws.cell(row=row, column=1, value="Provenance").font = _SECTION_FONT
    row += 1
    lines = [
        f"concentration policy: {provenance.get('concentration_policy')}",
        f"concentration fills: {provenance.get('fills')}",
        f"raw file: {provenance.get('raw_path')}",
        f"fill {data.get('fill_version')} · derive {data.get('derive_version')} · export {EXPORT_VERSION}",
    ]
    for line in lines:
        ws.cell(row=row, column=1, value=line).font = _MUTED_FONT
        row += 1

    notes = data.get("notes") or []
    if notes:
        ws.cell(row=row, column=1, value="Notes").font = _SECTION_FONT
        row += 1
        for note in notes:
            ws.cell(row=row, column=1, value=f"- {note}").font = _MUTED_FONT
            row += 1

    # ---- 列宽（便于阅读；数值列给足宽度）----
    widths = {1: 38, 2: 22, 3: 14, 4: 46}
    max_col = max(len(labels), 4)
    for index in range(1, max_col + 1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(index, 18)
    ws.sheet_view.showGridLines = True


def export_workbook(data: dict) -> Workbook:
    """把填充结果转成 openpyxl Workbook（每个离子一个工作表）。"""
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for sheet in data.get("sheets") or []:
        name = _safe_sheet_name(sheet.get("sheet_name") or sheet.get("analyte") or "Data", used)
        ws = workbook.create_sheet(name)
        _write_sheet(ws, data, sheet)
    if not workbook.sheetnames:
        ws = workbook.create_sheet("Data")
        ws["A1"] = "No data block available for template filling."
    return workbook


def save_workbook(data: dict, path: str | Path) -> Path:
    """写出 xlsx 到指定路径（父目录自动创建）。"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    export_workbook(data).save(target)
    return target


def workbook_bytes(data: dict) -> bytes:
    """在内存中生成 xlsx 字节流（供 API 直接下载）。"""
    buffer = io.BytesIO()
    export_workbook(data).save(buffer)
    return buffer.getvalue()
